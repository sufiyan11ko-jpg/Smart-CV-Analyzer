"""CV Analyser — password-only resume analysis and career matching app.
Run: python -m pip install -r requirements.txt && streamlit run app.py
"""
import os, re, io, json, hmac, hashlib, secrets, sqlite3
from datetime import datetime, timedelta
from xml.sax.saxutils import escape
from urllib.parse import quote_plus
import streamlit as st

st.set_page_config(page_title="CV Analyser Pro", page_icon="📄", layout="wide")
DB_FILE = os.getenv("CV_DB_FILE", "cv_analyser.db")
PBKDF2_ITERS = 310_000
MAX_OCR_PAGES = 10

# Broad starter catalogue: recommendations are inferred from resume skills, not live vacancies.
JOB_CATALOG = [
("Python Developer","IT & Software",["python","sql","git","rest api","django","flask"],"Build backend services, automation and Python applications."),
("Backend Developer","IT & Software",["python","java","node.js","sql","rest api","git","docker"],"Develop APIs, databases and server-side systems."),
("Frontend Developer","IT & Software",["html","css","javascript","react","git","ui/ux"],"Build responsive and accessible web interfaces."),
("Full Stack Developer","IT & Software",["html","css","javascript","react","node.js","sql","git","rest api"],"Work across frontend, backend and database layers."),
("Java Developer","IT & Software",["java","sql","git","rest api","spring"],"Develop enterprise and application software using Java."),
("Software Tester / QA Analyst","IT & Software",["testing","manual testing","automation testing","sql","python","jira"],"Test applications, report defects and improve quality."),
("IT Support Specialist","IT & Software",["linux","windows","networking","troubleshooting","hardware","communication"],"Troubleshoot user, device, network and software issues."),
("System Administrator","IT & Software",["linux","windows","networking","docker","aws","security"],"Maintain operating systems, servers and infrastructure."),
("Cloud Support Associate","Cloud & DevOps",["aws","linux","networking","python","docker","security"],"Support cloud infrastructure, deployments and incidents."),
("DevOps Engineer","Cloud & DevOps",["linux","docker","aws","git","python","ci/cd","kubernetes"],"Automate deployment, infrastructure and monitoring."),
("Cybersecurity Analyst","Cybersecurity",["networking","linux","security","python","siem","risk assessment"],"Monitor threats, investigate incidents and reduce security risks."),
("Data Analyst","Data & AI",["excel","sql","python","pandas","statistics","power bi","data visualization"],"Turn business data into reports, insights and decisions."),
("Business Intelligence Analyst","Data & AI",["sql","power bi","tableau","excel","data visualization","statistics"],"Create dashboards, KPIs and business reporting."),
("Junior Data Scientist","Data & AI",["python","pandas","numpy","statistics","machine learning","scikit-learn"],"Explore data and build predictive models."),
("Machine Learning Engineer","Data & AI",["python","machine learning","scikit-learn","tensorflow","pytorch","sql"],"Train, evaluate and deploy machine-learning systems."),
("AI / NLP Associate","Data & AI",["python","machine learning","nlp","pytorch","tensorflow","statistics"],"Work with language data, AI models and evaluation."),
("Data Engineer","Data & AI",["python","sql","pandas","aws","docker","etl","spark"],"Build data pipelines and reliable data platforms."),
("Database Administrator","Data & AI",["sql","mysql","mongodb","linux","backup","security"],"Maintain database performance, access and recovery."),
("Excel / MIS Executive","Business Operations",["excel","reporting","data analysis","communication","power bi"],"Prepare operational reports and maintain business data."),
("Financial Analyst","Finance & Accounting",["excel","financial analysis","forecasting","budgeting","accounting","power bi"],"Analyse financial performance, budgets and forecasts."),
("Accountant","Finance & Accounting",["accounting","bookkeeping","gst","taxation","tally","excel"],"Maintain accounts, reconciliations and statutory records."),
("Audit Associate","Finance & Accounting",["auditing","accounting","excel","taxation","risk assessment"],"Support audits, controls testing and financial checks."),
("HR Executive","Human Resources",["recruitment","onboarding","communication","excel","payroll","employee relations"],"Support hiring, employee records and HR operations."),
("Talent Acquisition Associate","Human Resources",["recruitment","talent acquisition","interviewing","communication","linkedin"],"Source candidates and coordinate hiring processes."),
("Digital Marketing Executive","Marketing & Sales",["seo","social media","content marketing","google analytics","google ads","copywriting"],"Plan and measure digital campaigns and content."),
("SEO Specialist","Marketing & Sales",["seo","sem","google analytics","content marketing","keyword research"],"Improve search visibility through technical and content SEO."),
("Social Media Executive","Marketing & Sales",["social media","content marketing","canva","copywriting","analytics"],"Create, schedule and analyse social media content."),
("Sales Development Representative","Marketing & Sales",["sales","crm","lead generation","communication","negotiation"],"Find prospects, qualify leads and build sales pipelines."),
("UI/UX Designer","Design & Creative",["figma","ui/ux","wireframing","prototyping","user research","typography"],"Research user needs and design usable digital products."),
("Graphic Designer","Design & Creative",["photoshop","illustrator","canva","typography","branding","graphic design"],"Create visual assets and brand communication."),
("Product Designer","Design & Creative",["figma","ui/ux","wireframing","prototyping","user research","design systems"],"Design end-to-end product experiences with teams."),
("Content Writer","Content & Communications",["copywriting","content marketing","research","communication","seo"],"Write clear content for websites, products and campaigns."),
("Technical Writer","Content & Communications",["technical writing","research","documentation","communication","api"],"Create product documentation, guides and tutorials."),
("Project Coordinator","Business Operations",["project management","excel","communication","reporting","time management"],"Coordinate timelines, stakeholders and project updates."),
("Operations Analyst","Business Operations",["excel","data analysis","reporting","process improvement","communication"],"Improve processes using operational data and reporting."),
("Customer Support Specialist","Customer Operations",["communication","problem solving","crm","troubleshooting","time management"],"Resolve customer issues and improve service experience."),
("Technical Support Engineer","Customer Operations",["troubleshooting","linux","networking","sql","python","communication"],"Diagnose technical problems and support users."),
("Network Engineer","Infrastructure",["networking","linux","security","troubleshooting","routing","switching"],"Configure, monitor and troubleshoot computer networks."),
("Embedded Systems Engineer","Engineering",["c++","c","embedded systems","iot","circuit design","linux"],"Develop firmware and software for connected devices."),
("Mechanical Design Engineer","Engineering",["autocad","solidworks","ansys","matlab","design","engineering"],"Design and validate mechanical components and systems."),
("Civil CAD Technician","Engineering",["autocad","revit","civil engineering","technical drawing","project management"],"Prepare technical drawings and construction documentation."),
("Research Assistant","Research & Education",["research","statistics","excel","documentation","communication"],"Support research, data collection and analysis."),
("Teaching / Training Assistant","Research & Education",["communication","presentation","teaching","training","time management"],"Support learning, training and educational delivery."),
("E-commerce Executive","Business Operations",["excel","inventory","sales","digital marketing","customer support"],"Manage online store operations, listings and orders."),
("Supply Chain Analyst","Business Operations",["excel","data analysis","forecasting","inventory","reporting"],"Analyse stock, logistics and supply chain performance."),
("Product Analyst","Data & AI",["sql","excel","data analysis","statistics","power bi","user research"],"Analyse product usage, funnels and customer behaviour."),
("Mobile App Developer","IT & Software",["java","kotlin","swift","javascript","react","api"],"Build and maintain Android or iOS applications."),
("Game Developer","IT & Software",["c++","c#","unity","python","git","3d"],"Create gameplay systems and interactive experiences."),
("Blockchain Developer","IT & Software",["solidity","javascript","python","cryptography","rest api","git"],"Build and test decentralised applications and smart contracts."),
]
SKILL_GROUPS = {
"IT & Software":["python","java","c++","c#","javascript","typescript","html","css","react","node.js","django","flask","streamlit","sql","mysql","mongodb","git","github","docker","linux","windows","rest api","api","testing","manual testing","automation testing","jira","spring","kotlin","swift","unity"],
"Data & AI":["python","sql","excel","pandas","numpy","scikit-learn","machine learning","tensorflow","pytorch","statistics","power bi","tableau","data analysis","data visualization","nlp","spark","etl"],
"Cloud & DevOps":["aws","azure","gcp","linux","docker","kubernetes","ci/cd","git","networking","terraform"],
"Cybersecurity":["security","networking","linux","siem","risk assessment","cryptography","penetration testing"],
"Finance & Accounting":["accounting","bookkeeping","financial analysis","budgeting","forecasting","auditing","taxation","gst","tally","quickbooks","sap","excel","payroll"],
"Marketing & Sales":["seo","sem","content marketing","social media","google analytics","email marketing","copywriting","branding","google ads","crm","lead generation","salesforce","sales","negotiation","keyword research"],
"Human Resources":["recruitment","talent acquisition","onboarding","employee relations","payroll","performance management","hris","interviewing"],
"Design & Creative":["figma","photoshop","illustrator","ui/ux","wireframing","prototyping","typography","graphic design","canva","user research","video editing","design systems"],
"Engineering":["autocad","solidworks","ansys","matlab","plc","scada","revit","circuit design","embedded systems","robotics","iot","quality control","c++","c"],
"General & Business":["communication","leadership","teamwork","problem solving","time management","critical thinking","adaptability","creativity","collaboration","presentation","project management","reporting","research","documentation","troubleshooting","process improvement","inventory","teaching","training","technical writing"],
}
SOFT = ["communication","leadership","teamwork","problem solving","time management","critical thinking","adaptability","creativity","collaboration","presentation"]
ACTION_VERBS = ["built","created","developed","designed","improved","increased","reduced","managed","led","implemented","analysed","analyzed","automated","delivered","launched","achieved","optimized","optimised","coordinated","tested"]


def connect():
    c=sqlite3.connect(DB_FILE,timeout=20); c.row_factory=sqlite3.Row; c.execute("PRAGMA foreign_keys=ON"); return c

def init_db():
    with connect() as c:
        c.executescript("""
        CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY AUTOINCREMENT,name TEXT NOT NULL,email TEXT UNIQUE,mobile TEXT UNIQUE,password_hash TEXT NOT NULL,created_at TEXT NOT NULL,failed_attempts INTEGER DEFAULT 0,locked_until TEXT);
        CREATE TABLE IF NOT EXISTS resume_history(id INTEGER PRIMARY KEY AUTOINCREMENT,user_id INTEGER NOT NULL,file_name TEXT NOT NULL,score INTEGER NOT NULL,industry TEXT,job_match INTEGER,analysis_json TEXT NOT NULL,extracted_text TEXT,created_at TEXT NOT NULL,FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE);
        CREATE INDEX IF NOT EXISTS idx_history_user ON resume_history(user_id);
        """)
init_db()

def now(): return datetime.now()
def stamp(dt=None): return (dt or now()).strftime("%Y-%m-%d %H:%M:%S")
def hash_password(pw):
    salt=secrets.token_bytes(16); digest=hashlib.pbkdf2_hmac("sha256",pw.encode(),salt,PBKDF2_ITERS); return f"{PBKDF2_ITERS}${salt.hex()}${digest.hex()}"
def verify_password(pw,stored):
    try:
        rounds,salt,digest=stored.split("$"); got=hashlib.pbkdf2_hmac("sha256",pw.encode(),bytes.fromhex(salt),int(rounds)).hex(); return hmac.compare_digest(got,digest)
    except Exception: return False
def valid_password(p): return len(p)>=8 and bool(re.search(r"[A-Za-z]",p)) and bool(re.search(r"\d",p))
def normalize_mobile(v):
    d=re.sub(r"\D","",v or ""); return d[-10:] if len(d) in (10,11,12,13,14,15) else None
def identify(v):
    v=(v or "").strip()
    if "@" in v: return ("email",v.lower()) if re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+",v) else (None,None)
    m=normalize_mobile(v); return ("mobile",m) if m else (None,None)
def get_user(kind,value):
    if not kind:return None
    with connect() as c:return c.execute(f"SELECT * FROM users WHERE {kind}=?",(value,)).fetchone()

def ocr_image(img):
    import pytesseract
    from PIL import ImageOps
    img=ImageOps.autocontrast(img.convert("L"))
    if max(img.size)<1800:
        scale=1800/max(img.size); img=img.resize((int(img.width*scale),int(img.height*scale)))
    return pytesseract.image_to_string(img,config="--psm 3")
def extract_file(upload):
    name,data=upload.name,upload.getvalue(); lower=name.lower()
    if data.startswith(b"%PDF"):
        import pdfplumber
        texts=[]
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for i,page in enumerate(pdf.pages):
                txt=page.extract_text() or ""
                if len(txt.strip())<25 and i<MAX_OCR_PAGES:
                    try:txt=ocr_image(page.to_image(resolution=200).original)
                    except Exception:pass
                texts.append(txt)
        return "\n".join(texts),name
    if lower.endswith(".docx"):
        from docx import Document
        doc=Document(io.BytesIO(data)); return "\n".join([p.text for p in doc.paragraphs]+[" | ".join(x.text for x in r.cells) for t in doc.tables for r in t.rows]),name
    if lower.endswith((".png",".jpg",".jpeg",".webp")):
        from PIL import Image
        return ocr_image(Image.open(io.BytesIO(data))),name
    if lower.endswith(".txt"):return data.decode("utf-8",errors="ignore"),name
    raise ValueError("Supported formats: PDF, DOCX, TXT, JPG, PNG and WEBP.")

def has_term(text,term):
    return bool(re.search(r"(?<![\w+#])"+re.escape(term.lower())+r"(?![\w+#])",text.lower()))
def detect_skills(text):
    return sorted({s for group in SKILL_GROUPS.values() for s in group if has_term(text,s)})
def section_lines(lines,aliases):
    sections={}; current=None
    for line in lines:
        key=re.sub(r"[^a-z ]","",line.lower()).strip()
        match=next((k for k,vals in aliases.items() if key in vals),None)
        if match:current=match;sections.setdefault(current,[]);continue
        if current:sections.setdefault(current,[]).append(line)
    return sections

def analyse(text,filename,jd=""):
    lines=[x.strip(" •\t-") for x in text.splitlines() if x.strip()]
    email=re.search(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}",text)
    phone=re.search(r"(?:\+?\d[\d\s().-]{8,}\d)",text)
    name="Not found"
    for line in lines[:12]:
        if "@" not in line and not re.search(r"\d",line) and 2<=len(line.split())<=4 and len(line)<45:name=line;break
    aliases={"Education":["education","academic qualification","qualifications"],"Experience":["experience","work experience","employment","internship"],"Projects":["projects","project"],"Certifications":["certifications","certification","certificates","courses"],"Summary":["summary","professional summary","profile","objective"],"Skills":["skills","technical skills","key skills"]}
    sections=section_lines(lines,aliases)
    education=sections.get("Education",[])[:15]; experience=sections.get("Experience",[])[:20]; projects=sections.get("Projects",[])[:15]; certs=sections.get("Certifications",[])[:12]
    skills=detect_skills(text); soft=[s for s in SOFT if has_term(text,s)]; words=len(text.split())
    bullets=[x for x in lines if re.match(r"^[•*▪\-]",x) or len(x.split())>=8]
    quantified=sum(bool(re.search(r"\b\d+(?:\.\d+)?\s*(?:%|percent|users|customers|hours|days|months|years|x|k|m)\b",x,re.I)) for x in bullets)
    action_count=sum(any(re.search(r"\b"+v+r"\b",x,re.I) for v in ACTION_VERBS) for x in bullets)
    links=re.findall(r"(?:https?://)?(?:www\.)?(?:linkedin\.com|github\.com)/[^\s)]+",text,re.I)
    # Transparent heuristic score, not an ATS vendor score.
    checks={
      "Contact details":15 if email and phone else 10 if email or phone else 0,
      "Core sections":min(20, (5 if education else 0)+(6 if experience else 0)+(5 if projects else 0)+(4 if sections.get("Summary") else 0)),
      "Relevant skills":min(20,len(skills)*2),
      "Evidence & impact":min(20,(5 if experience else 0)+(5 if projects else 0)+min(5,action_count)+(5 if quantified else 0)),
      "ATS readability":10 if words>=120 and len(re.findall(r"[^\x00-\x7F]",text))<max(15,words*.2) else 6 if words>=60 else 2,
      "Length & completeness":15 if 250<=words<=900 else 10 if 120<=words<250 or 900<words<=1200 else 4,
    }
    score=round(sum(checks.values()))
    role_results=[]
    for role,sector,required,desc in JOB_CATALOG:
        matched=[s for s in required if s in skills]
        missing=[s for s in required if s not in skills]
        # Match is a skills coverage estimate; no claim of hiring probability.
        match=round(100*len(matched)/len(required)) if required else 0
        role_results.append({"role":role,"sector":sector,"description":desc,"match":match,"matched":matched,"missing":missing,"required_count":len(required)})
    role_results.sort(key=lambda x:(x["match"],len(x["matched"])),reverse=True)
    jdskills=detect_skills(jd) if jd.strip() else []; jdmatched=[s for s in jdskills if s in skills]
    jdmatch=round(100*len(jdmatched)/len(jdskills)) if jdskills else 0
    tips=[]
    if not email or not phone:tips.append("Add a professional email address and a reachable phone number near the top.")
    if not sections.get("Summary"):tips.append("Add a 2–3 line professional summary tailored to your target role.")
    if not education:tips.append("Add an Education section with qualification, institute and graduation year.")
    if not experience:tips.append("If you are a fresher, add internships, volunteering, freelance work or practical training instead of inventing experience.")
    if len(projects)<2:tips.append("Add 2–3 relevant projects with your role, tools, problem and outcome.")
    if len(skills)<8:tips.append("Add relevant skills that you genuinely know; avoid keyword stuffing.")
    if action_count<3:tips.append("Start bullet points with action verbs such as Built, Analysed, Designed or Improved.")
    if quantified==0:tips.append("Add truthful measurable outcomes where possible (%, time saved, users, records or accuracy).")
    if not links:tips.append("Consider adding a LinkedIn or GitHub link if it supports your target role.")
    if words<120:tips.append("Resume text looks short; include relevant projects, skills and education details.")
    if words>1200:tips.append("Consider shortening the resume and prioritising the most relevant information.")
    return {"file_name":filename,"created":stamp(),"contact":{"name":name,"email":email.group(0) if email else "Not found","phone":phone.group(0).strip() if phone else "Not found"},"skills":skills,"soft_skills":soft,"industry":role_results[0]["sector"] if role_results and role_results[0]["match"]>0 else "Not determined yet","education":education,"experience":experience,"projects":projects,"certifications":certs,"summary":sections.get("Summary",[])[:6],"score":score,"score_breakdown":checks,"job_match":jdmatch,"jd_skills":jdskills,"jd_matched":jdmatched,"roles":role_results[:15],"role_count":len(JOB_CATALOG),"tips":tips,"words":words,"bullets":len(bullets),"action_bullets":action_count,"quantified_bullets":quantified,"links":links,"soft_count":len(soft)}

def make_pdf(a):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    b=io.BytesIO();styles=getSampleStyleSheet();story=[Paragraph("CV Analyser Pro — Resume Report",styles["Title"]),Spacer(1,10),Paragraph(f"File: {escape(a['file_name'])} | Date: {a['created']}",styles["BodyText"]),Paragraph(f"Resume score: {a['score']}/100 | Suggested field: {escape(a['industry'])}",styles["Heading2"]),Spacer(1,8)]
    sections=[("Score breakdown",list(a["score_breakdown"].items())),("Candidate",list(a["contact"].items())),("Skills",[("Detected",", ".join(a["skills"]) or "None"),("Soft skills",", ".join(a["soft_skills"]) or "None")]),("Top career matches",[(f"{r['role']} ({r['match']}%)",f"Matched: {', '.join(r['matched']) or 'None'} | Learn: {', '.join(r['missing']) or '—'}") for r in a["roles"][:8]]),("Education",[(str(i+1),v) for i,v in enumerate(a["education"])]),("Experience",[(str(i+1),v) for i,v in enumerate(a["experience"])]),("Projects",[(str(i+1),v) for i,v in enumerate(a["projects"])]),("Certifications",[(str(i+1),v) for i,v in enumerate(a["certifications"])]),("Improvements",[(str(i+1),v) for i,v in enumerate(a["tips"])])]
    for heading,vals in sections:
        story.append(Paragraph(heading,styles["Heading2"]));rows=[[Paragraph(escape(str(k)),styles["BodyText"]),Paragraph(escape(str(v)),styles["BodyText"])] for k,v in vals] or [["—","Not found"]];t=Table(rows,colWidths=[4.2*cm,12.8*cm]);t.setStyle(TableStyle([("GRID",(0,0),(-1,-1),.4,colors.lightgrey),("VALIGN",(0,0),(-1,-1),"TOP"),("BACKGROUND",(0,0),(0,-1),colors.whitesmoke)]));story.extend([t,Spacer(1,8)])
    story.append(Paragraph("Note: This is a heuristic resume review and skill-coverage estimate, not a guaranteed ATS score, hiring decision or live vacancy check.",styles["Italic"]))
    SimpleDocTemplate(b,pagesize=A4).build(story);return b.getvalue()

def save_result(uid,a,text):
    with connect() as c:c.execute("INSERT INTO resume_history(user_id,file_name,score,industry,job_match,analysis_json,extracted_text,created_at) VALUES(?,?,?,?,?,?,?,?)",(uid,a["file_name"],a["score"],a["industry"],a["job_match"],json.dumps(a),text,stamp()))

def style():
    st.markdown("""<style>.stApp{background:#f5f7fb;color:#182033}[data-testid="stSidebar"]{background:#fff}.hero{padding:24px;border-radius:18px;background:linear-gradient(120deg,#3347a9,#078b88);color:white;margin-bottom:18px}.hero h1,.hero p{color:white!important}div[data-testid="stMetric"]{background:white;padding:14px;border:1px solid #e5e7eb;border-radius:14px}</style>""",unsafe_allow_html=True)
style()
for k,v in {"logged_in":False,"uid":None,"user_name":"","result":None}.items():st.session_state.setdefault(k,v)

def login_ui():
    st.markdown('<div class="hero"><h1>📄 CV Analyser Pro</h1><p>Resume quality review · Skill analysis · Career matching · Reports</p></div>',unsafe_allow_html=True)
    t1,t2=st.tabs(["🔐 Login","📝 Sign Up"])
    with t1:
        with st.form("login"):
            ident=st.text_input("Email or mobile");pw=st.text_input("Password",type="password");go=st.form_submit_button("Login")
        if go:
            kind,val=identify(ident);u=get_user(kind,val)
            if not u:st.error("Invalid credentials.")
            elif u["locked_until"] and u["locked_until"]>stamp():st.error("Account temporarily locked. Try again later.")
            elif not verify_password(pw,u["password_hash"]):
                with connect() as c:
                    fails=u["failed_attempts"]+1;c.execute("UPDATE users SET failed_attempts=?,locked_until=? WHERE id=?",(0 if fails>=5 else fails,stamp(now()+timedelta(minutes=10)) if fails>=5 else None,u["id"]))
                st.error("Invalid credentials.")
            else:
                with connect() as c:c.execute("UPDATE users SET failed_attempts=0,locked_until=NULL WHERE id=?",(u["id"],))
                st.session_state.update(logged_in=True,uid=u["id"],user_name=u["name"]);st.rerun()
    with t2:
        with st.form("signup"):
            name=st.text_input("Full name");email=st.text_input("Email (optional if mobile supplied)");mobile=st.text_input("Mobile (optional if email supplied)");p1=st.text_input("Password (8+ characters, letters and numbers)",type="password");p2=st.text_input("Confirm password",type="password");go=st.form_submit_button("Create account")
        if go:
            em=email.strip().lower() or None;mob=normalize_mobile(mobile) if mobile.strip() else None
            if not name.strip():st.error("Enter your name.")
            elif not em and not mob:st.error("Enter email or mobile.")
            elif em and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+",em):st.error("Invalid email.")
            elif mobile.strip() and not mob:st.error("Invalid mobile number.")
            elif not valid_password(p1):st.error("Password must be 8+ characters with letters and numbers.")
            elif p1!=p2:st.error("Passwords do not match.")
            else:
                try:
                    with connect() as c:c.execute("INSERT INTO users(name,email,mobile,password_hash,created_at) VALUES(?,?,?,?,?)",(name.strip(),em,mob,hash_password(p1),stamp()))
                    st.success("Account created. Please log in.")
                except sqlite3.IntegrityError:st.error("Email or mobile already registered.")

def role_search_link(role):return "https://www.google.com/search?q="+quote_plus(role+" jobs worldwide")
def dashboard(a):
    st.subheader("📌 Resume overview")
    c1,c2,c3,c4=st.columns(4);c1.metric("Resume score",f"{a['score']}/100");c2.metric("Skills detected",len(a["skills"]));c3.metric("Career roles checked",a["role_count"]);c4.metric("Word count",a["words"])
    st.progress(a["score"]/100)
    st.caption("Score is a transparent rule-based review, not a score issued by an employer or commercial ATS.")
    left,right=st.columns(2)
    with left:
        st.markdown("#### Resume quality breakdown");st.bar_chart({"Points earned":a["score_breakdown"]})
    with right:
        st.markdown("#### Resume writing signals");st.bar_chart({"Count":{"Bullets":a["bullets"],"Action-led bullets":a["action_bullets"],"Bullets with numbers":a["quantified_bullets"],"Soft skills":a["soft_count"]}})
    st.subheader("👤 Extracted candidate details");st.json(a["contact"])
    tabs=st.tabs(["Skills & ATS","Best career matches","Resume sections","Action plan"])
    with tabs[0]:
        st.markdown("#### Detected skills");st.write(", ".join(a["skills"]) or "No known skills detected.")
        st.markdown("#### Skill distribution by category")
        counts={g:sum(1 for s in a["skills"] if s in ss) for g,ss in SKILL_GROUPS.items()};counts={k:v for k,v in counts.items() if v}
        if counts:st.bar_chart(counts)
        else:st.info("No catalogue skills found. Add skills you genuinely have, or check the extracted text.")
        st.markdown("#### ATS-style checklist")
        checks=[("Email and phone detected",a["contact"]["email"]!="Not found" and a["contact"]["phone"]!="Not found"),("Professional summary",bool(a["summary"])),("Education section",bool(a["education"])),("Experience or internship",bool(a["experience"])),("Projects section",bool(a["projects"])),("Certifications section",bool(a["certifications"])),("LinkedIn or GitHub link",bool(a["links"])),("Action verbs in bullet points",a["action_bullets"]>=3),("Quantified achievements",a["quantified_bullets"]>0)]
        for label,ok in checks:st.write(("✅ " if ok else "⚠️ ")+label)
    with tabs[1]:
        st.markdown(f"#### {a['role_count']} career paths compared using resume skills")
        st.caption("Match % = share of listed role skills found in this resume. It is not a probability of getting hired and does not check live openings.")
        top=a["roles"][:10]
        if top:
            st.bar_chart({"Skill coverage %":{r["role"]:r["match"] for r in top}})
            for i,r in enumerate(top,1):
                with st.expander(f"{i}. {r['role']} — {r['match']}% skill coverage | {r['sector']}"):
                    st.write(r["description"]);st.write("**Skills found:**",", ".join(r["matched"]) or "None yet");st.write("**Skills to develop:**",", ".join(r["missing"]) or "No gaps in this starter checklist")
                    st.markdown(f"[Search current {r['role']} vacancies worldwide]({role_search_link(r['role'])})")
        if a["jd_skills"]:
            st.markdown("#### Your pasted job description match");st.metric("Skill overlap",f"{a['job_match']}%");st.write("Matched:",", ".join(a["jd_matched"]) or "None");st.write("Missing:",", ".join([s for s in a["jd_skills"] if s not in a["jd_matched"]]) or "None")
    with tabs[2]:
        for key in ("summary","education","experience","projects","certifications"):
            with st.expander(key.replace("_"," ").title()):st.write(a[key] or "Not found in extracted resume text.")
    with tabs[3]:
        st.markdown("#### Prioritised improvement plan")
        for i,tip in enumerate(a["tips"],1):st.write(f"{i}. {tip}")
        st.markdown("#### Suggested next steps")
        st.write("1. Pick one target role from the career matches.")
        st.write("2. Tailor the summary, skills and project bullets to that role without adding false claims.")
        st.write("3. Compare the revised CV with a real job description before applying.")
    try:
        pdf=make_pdf(a);st.download_button("⬇️ Download detailed PDF report",pdf,file_name="CV_Analyser_Pro_Report.pdf",mime="application/pdf")
    except Exception as e:st.warning(f"PDF report unavailable: {e}")
    st.download_button("⬇️ Download analysis JSON",json.dumps(a,indent=2),file_name="CV_Analyser_Analysis.json",mime="application/json")

def app_ui():
    with st.sidebar:
        st.markdown(f"### Welcome, {st.session_state.user_name}");page=st.radio("Navigation",["Analyze Resume","Dashboard","History","Profile"])
        if st.button("Logout"):st.session_state.update(logged_in=False,uid=None,user_name="",result=None);st.rerun()
    if page=="Analyze Resume":
        st.title("🔍 Professional Resume Analysis")
        method=st.radio("Input type",["Upload File","Paste Text"],horizontal=True);uploaded=None;pasted=""
        if method=="Upload File":uploaded=st.file_uploader("Upload resume",type=["pdf","docx","txt","jpg","jpeg","png","webp"])
        else:pasted=st.text_area("Paste resume text",height=240)
        jd=st.text_area("Paste a target job description (optional) for direct skill matching",height=130)
        if st.button("🚀 Run Professional Resume Check",type="primary"):
            try:
                if method=="Upload File":
                    if not uploaded:st.error("Upload a file first.");st.stop()
                    with st.spinner("Reading resume and checking sections, skills and career paths..."):text,filename=extract_file(uploaded)
                else:text,filename=pasted,"Pasted_Resume.txt"
                if len(text.strip())<30:st.error("Not enough readable text. Try a clearer scan or paste text.")
                else:
                    a=analyse(text,filename,jd);save_result(st.session_state.uid,a,text);st.session_state.result={"analysis":a,"text":text}
            except Exception as e:st.error(f"Could not analyse file: {e}")
        if st.session_state.result:
            a=st.session_state.result["analysis"];dashboard(a)
            with st.expander("View extracted resume text"):st.text(st.session_state.result["text"][:15000])
    elif page=="Dashboard":
        st.title("📊 Analysis Dashboard")
        with connect() as c:rows=c.execute("SELECT score,job_match,created_at FROM resume_history WHERE user_id=? ORDER BY id",(st.session_state.uid,)).fetchall()
        if rows:
            x,y,z=st.columns(3);x.metric("Total analyses",len(rows));y.metric("Average score",round(sum(r["score"] for r in rows)/len(rows)));z.metric("Best score",max(r["score"] for r in rows));st.line_chart({"Resume score":[r["score"] for r in rows]})
        else:st.info("No analyses yet. Start from Analyze Resume.")
    elif page=="History":
        st.title("📚 Resume History")
        with connect() as c:rows=c.execute("SELECT * FROM resume_history WHERE user_id=? ORDER BY id DESC",(st.session_state.uid,)).fetchall()
        if not rows:st.info("No saved resume history.")
        for row in rows:
            with st.expander(f"{row['file_name']} · Score {row['score']}/100 · {row['created_at']}"):
                a=json.loads(row["analysis_json"]);dashboard(a)
                with st.expander("Extracted text"):st.text((row["extracted_text"] or "")[:10000])
                if st.button("Delete this record",key=f"del_{row['id']}"):
                    with connect() as c:c.execute("DELETE FROM resume_history WHERE id=? AND user_id=?",(row["id"],st.session_state.uid))
                    st.rerun()
        if rows and st.button("Clear all history"):
            with connect() as c:c.execute("DELETE FROM resume_history WHERE user_id=?",(st.session_state.uid,))
            st.rerun()
    else:
        st.title("👤 Profile")
        with connect() as c:u=c.execute("SELECT * FROM users WHERE id=?",(st.session_state.uid,)).fetchone()
        st.write({"Name":u["name"],"Email":u["email"] or "—","Mobile":u["mobile"] or "—","Member since":u["created_at"]})
        with st.form("change_pw"):
            old=st.text_input("Current password",type="password");new=st.text_input("New password",type="password");confirm=st.text_input("Confirm new password",type="password");go=st.form_submit_button("Change password")
        if go:
            if not verify_password(old,u["password_hash"]):st.error("Current password incorrect.")
            elif not valid_password(new):st.error("New password must be 8+ characters with letters and numbers.")
            elif new!=confirm:st.error("Passwords do not match.")
            else:
                with connect() as c:c.execute("UPDATE users SET password_hash=? WHERE id=?",(hash_password(new),u["id"]))
                st.success("Password changed.")

st.markdown('<div class="hero"><h1>📄 CV Analyser Pro</h1><p>Resume quality review · ATS-style checks · Global career matching · PDF reports</p></div>',unsafe_allow_html=True)
if not st.session_state.logged_in:login_ui()
else:app_ui()
